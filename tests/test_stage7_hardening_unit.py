import datetime as dt
import uuid
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from sqlalchemy.exc import InvalidRequestError

from tests.db_safety import validate_test_database_url
from app.core.exceptions import AccessDeniedError, ValidationError, ConflictError
from app.core.numbers import volume, money
from app.core.operations import operation, request_id
from app.bot.keyboards.act import ActWorkCallback, ActLineCallback, ActCardActionCallback, ActPageCallback
from app.reports import act_excel
from tests.test_stage7_excel import _make_act_with_lines
from app.services import act, customer_acceptance as ca, access, audit


@pytest.mark.parametrize('target,app', [
    (None, None), ('postgresql+asyncpg://u:p@db/work_accounting', None),
    ('postgresql+asyncpg://u:p@db/demo', None),
    ('postgresql+asyncpg://u:p@db/app_test', 'postgresql+asyncpg://u:p@db/app_test'),
    ('postgresql+asyncpg://u:other@db/app_test', 'postgresql+asyncpg://v:p@db:5432/app_test'),
    ('postgresql+asyncpg://u:p@db/app_test?database=work_accounting', None),
    ('sqlite:///app_test', None), ('not a url with SECRET', None),
])
def test_reject_unsafe_test_targets(target, app):
    with pytest.raises(ValueError) as error:
        validate_test_database_url(target, app)
    assert 'SECRET' not in str(error.value)


def test_accept_explicit_test_target():
    assert validate_test_database_url('postgresql+asyncpg://u:p@db/work_accounting_test',
                                      'postgresql+asyncpg://u:p@db/work_accounting').database == 'work_accounting_test'


@pytest.mark.parametrize('value', ['NaN', 'sNaN', 'Infinity', '-Infinity', '0.0001', '1000000000000000'])
def test_volume_rejects_nonfinite_precision_and_overflow(value):
    with pytest.raises(ValidationError):
        volume(Decimal(value))


@pytest.mark.parametrize('value', ['NaN', 'Infinity', '1.001', '10000000000000000'])
def test_money_rejects_nonfinite_precision_and_overflow(value):
    with pytest.raises(ValidationError):
        money(Decimal(value))


def test_numeric_boundary():
    assert volume(Decimal('999999999999999.999')) == Decimal('999999999999999.999')
    assert money(Decimal('9999999999999999.99')) == Decimal('9999999999999999.99')


@pytest.mark.parametrize('callback', [ActWorkCallback(work_entry_id=uuid.uuid4()),
    ActLineCallback(line_id=uuid.uuid4()), ActLineCallback(line_id=uuid.uuid4(), action='cancel'),
    ActCardActionCallback(act_id=uuid.uuid4(), action='add_line'), ActPageCallback(kind='contracts', page=99999)])
def test_actual_callback_pack_below_64_bytes(callback):
    packed = callback.pack()
    assert len(packed.encode()) < 64
    assert type(callback).unpack(packed) == callback


@pytest.mark.parametrize('allowed', [None, uuid.uuid4()])
async def test_permission_is_checked_on_server(allowed):
    session = SimpleNamespace(scalar=AsyncMock(return_value=allowed))
    if allowed is None:
        with pytest.raises(AccessDeniedError):
            await access.require_permission(session, uuid.uuid4(), 'acts.manage')
    else:
        await access.require_permission(session, uuid.uuid4(), 'acts.manage')
    session.scalar.assert_awaited_once()


async def test_anonymous_user_denied_without_query():
    session = SimpleNamespace(scalar=AsyncMock())
    with pytest.raises(AccessDeniedError):
        await access.require_permission(session, None, 'acts.view')
    session.scalar.assert_not_awaited()


async def test_operation_correlation_shared_and_reset():
    seen = []
    @operation
    async def nested():
        seen.append(request_id.get())
    @operation
    async def outer():
        seen.append(request_id.get())
        await nested()
    await outer()
    await outer()
    assert seen[0] == seen[1] and seen[2] == seen[3] and seen[0] != seen[2]
    assert request_id.get() is None


def test_excel_values_remain_decimal_and_text_is_not_formula(tmp_path):
    values = _make_act_with_lines()
    values[1][0].comment = '=HYPERLINK("https://example.invalid","click")'
    values[1][0].work_type_name = '=1+1'
    workbook = act_excel.build_act_workbook(*values)
    assert isinstance(workbook.active['I5'].value, Decimal)
    assert workbook.active['J5'].data_type == 's'
    assert workbook.active['D5'].data_type == 's'
    target = tmp_path / 'test.xlsx'
    workbook.save(target)
    from openpyxl import load_workbook
    loaded = load_workbook(target)
    assert loaded.active['J5'].data_type == 's'
    assert loaded.active['D5'].value == '=1+1'
    loaded.close()
    workbook.close()


def test_unique_excel_paths_and_cleanup_after_delivery_failure():
    values = _make_act_with_lines()
    with pytest.raises(RuntimeError):
        with act_excel.temporary_act_excel(*values) as first:
            with act_excel.temporary_act_excel(*values) as second:
                assert first != second and first.exists() and second.exists()
                raise RuntimeError('delivery failed')
    assert not first.exists() and not second.exists()


def test_excel_cleanup_after_partial_save_failure(monkeypatch, tmp_path):
    def bad_save(self, filename):
        Path(filename).write_text('partial')
        raise OSError('disk error')
    monkeypatch.setattr(act_excel.tempfile, 'tempdir', str(tmp_path))
    monkeypatch.setattr(act_excel.Workbook, 'save', bad_save)
    with pytest.raises(OSError):
        act_excel.generate_act_excel(*_make_act_with_lines())
    assert list(tmp_path.iterdir()) == []


async def test_choose_contract_begins_before_first_query(monkeypatch):
    from app.bot.handlers import act as handler
    class Session:
        active = False
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
        def begin(self):
            if self.active: raise InvalidRequestError('already begun')
            self.active = True
            return self
        async def get(self, *args):
            assert self.active, 'Query must be inside the transaction'
            return SimpleNamespace(code='P')
    monkeypatch.setattr(handler, 'session_factory', Session)
    create = AsyncMock(return_value=SimpleNamespace(id=uuid.uuid4()))
    monkeypatch.setattr(handler.act_service, 'create_draft', create)
    monkeypatch.setattr(handler, '_show_act_card', AsyncMock())
    state = SimpleNamespace(get_data=AsyncMock(return_value={'project_id': str(uuid.uuid4())}))
    callback = SimpleNamespace(answer=AsyncMock())
    await handler.choose_contract(callback, SimpleNamespace(contract_id='none'), state, SimpleNamespace(id=uuid.uuid4()))
    create.assert_awaited_once()


@pytest.fixture
def service_mocks(monkeypatch):
    session = SimpleNamespace(scalar=AsyncMock(return_value=None), get=AsyncMock(),
        flush=AsyncMock(), add=MagicMock(), delete=AsyncMock())
    for module in (act, ca):
        monkeypatch.setattr(module, 'require_permission', AsyncMock())
        monkeypatch.setattr(module, 'lock_key', AsyncMock())
    return session


async def test_acceptance_replay_returns_existing_before_volume_consumed(service_mocks):
    existing = SimpleNamespace(work_entry_id=uuid.uuid4(), accepted_volume=Decimal('10'))
    service_mocks.scalar.return_value = existing
    result = await ca.accept(service_mocks, work_entry_id=existing.work_entry_id, volume=Decimal('10'),
                             idempotency_key=uuid.uuid4(), user_id=uuid.uuid4())
    assert result is existing
    service_mocks.add.assert_not_called()


async def test_line_replay_returns_existing_even_if_finalized(service_mocks):
    existing = SimpleNamespace(act_id=uuid.uuid4(), work_entry_id=uuid.uuid4(), accepted_volume=Decimal('10'))
    service_mocks.scalar.return_value = existing
    result = await act.add_line(service_mocks, act_id=existing.act_id, work_entry_id=existing.work_entry_id,
        quantity=Decimal('10'), idempotency_key=uuid.uuid4(), user_id=uuid.uuid4())
    assert result is existing
    service_mocks.add.assert_not_called()


async def test_line_key_cannot_be_reused_for_different_payload(service_mocks):
    service_mocks.scalar.return_value = SimpleNamespace(act_id=uuid.uuid4())
    with pytest.raises(ConflictError):
        await act.add_line(service_mocks, act_id=uuid.uuid4(), work_entry_id=uuid.uuid4(),
                           quantity=Decimal('10'), idempotency_key=uuid.uuid4(), user_id=uuid.uuid4())


async def test_correction_cannot_reduce_below_active_acts(service_mocks, monkeypatch):
    monkeypatch.setattr(ca, 'lock_work_entry', AsyncMock(return_value=SimpleNamespace(status='included_in_act')))
    monkeypatch.setattr(ca.ca_repo, 'total_accepted_volume', AsyncMock(return_value=Decimal('100')))
    monkeypatch.setattr(ca.act_repo, 'volume_in_active_acts', AsyncMock(return_value=Decimal('90')))
    with pytest.raises(ValidationError, match='действующих актов'):
        await ca.correct(service_mocks, work_entry_id=uuid.uuid4(), delta_volume=Decimal('-20'), user_id=uuid.uuid4(), reason='fix')
    service_mocks.add.assert_not_called()


@pytest.mark.parametrize('status,mismatch', [('cancelled', False), ('accepted_customer', True)])
async def test_reject_wrong_project_and_cancelled_work(service_mocks, monkeypatch, status, mismatch):
    aid, pid, cid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    monkeypatch.setattr(act.act_repo, 'get_for_update', AsyncMock(return_value=SimpleNamespace(id=aid, status='draft', project_id=pid, contract_id=cid)))
    monkeypatch.setattr(act, 'lock_work_entry', AsyncMock(return_value=SimpleNamespace(status=status,
                        project_id=uuid.uuid4() if mismatch else pid, contract_id=cid)))
    with pytest.raises(ValidationError):
        await act.add_line(service_mocks, act_id=aid, work_entry_id=uuid.uuid4(), quantity=Decimal('1'),
                           idempotency_key=uuid.uuid4(), user_id=uuid.uuid4())
    service_mocks.add.assert_not_called()


async def test_draft_contract_must_belong_to_project(service_mocks):
    service_mocks.get.side_effect = [SimpleNamespace(is_archived=False), SimpleNamespace(project_id=uuid.uuid4())]
    with pytest.raises(ValidationError, match='Договор'):
        await act.create_draft(service_mocks, project_id=uuid.uuid4(), contract_id=uuid.uuid4(), act_number='x',
                               act_date=dt.date.today(), created_by=uuid.uuid4())


async def test_cancel_line_preserves_record_and_metadata(service_mocks, monkeypatch):
    aid, wid, uid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    line = SimpleNamespace(id=uuid.uuid4(), act_id=aid, work_entry_id=wid, cancelled_at=None, accepted_volume=Decimal('1'), amount=Decimal('2'))
    monkeypatch.setattr(act.act_repo, 'get_for_update', AsyncMock(return_value=SimpleNamespace(id=aid, status='draft')))
    monkeypatch.setattr(act.act_repo, 'get_line', AsyncMock(return_value=line))
    monkeypatch.setattr(act, 'lock_work_entry', AsyncMock(return_value=SimpleNamespace(id=wid)))
    monkeypatch.setattr(act, '_recalculate_total', AsyncMock())
    sync = AsyncMock()
    monkeypatch.setattr(act, 'synchronize', sync)
    monkeypatch.setattr(act.audit, 'record', AsyncMock())
    await act.remove_line(service_mocks, act_id=aid, line_id=line.id, reason='mistake', user_id=uid)
    assert line.cancelled_by == uid and line.cancelled_at.tzinfo is not None
    assert line.cancellation_reason == 'mistake'
    service_mocks.delete.assert_not_awaited()
    sync.assert_awaited_once()


async def test_status_change_records_history(monkeypatch):
    from app.services import work_status
    monkeypatch.setattr(work_status.act_repo, 'volume_in_active_acts', AsyncMock(return_value=Decimal('0')))
    monkeypatch.setattr(work_status.ca_repo, 'total_accepted_volume', AsyncMock(return_value=Decimal('100')))
    monkeypatch.setattr(work_status, 'total_internal_accepted_volume', AsyncMock(return_value=Decimal('100')))
    work = SimpleNamespace(id=uuid.uuid4(), status='included_in_act')
    session = SimpleNamespace(add=MagicMock())
    await work_status.synchronize(session, work, uuid.uuid4(), 'cancel')
    assert work.status == 'accepted_customer'
    history = session.add.call_args.args[0]
    assert history.old_status == 'included_in_act' and history.new_status == 'accepted_customer'


async def test_export_audited(service_mocks, monkeypatch):
    aid = uuid.uuid4()
    record = SimpleNamespace(id=aid, project_id=uuid.uuid4(), contract_id=None)
    monkeypatch.setattr(act.act_repo, 'get_for_update', AsyncMock(return_value=record))
    monkeypatch.setattr(act.act_repo, 'list_lines', AsyncMock(return_value=[]))
    audit_mock = AsyncMock()
    monkeypatch.setattr(act.audit, 'record', audit_mock)
    await act.prepare_export(service_mocks, act_id=aid, user_id=uuid.uuid4())
    assert audit_mock.call_args.kwargs['action'] == 'act.export'


def test_excel_preserves_large_exact_money(tmp_path):
    values = _make_act_with_lines()
    values[0].total_amount = Decimal('9999999999999999.99')
    workbook = act_excel.build_act_workbook(*values)
    assert workbook.active['I7'].value == '9999999999999999.99'
    target = tmp_path / 'large.xlsx'
    workbook.save(target)
    from openpyxl import load_workbook
    loaded = load_workbook(target)
    assert loaded.active['I7'].value == '9999999999999999.99'
    loaded.close()
    workbook.close()


async def test_permission_middleware_denies_handler(monkeypatch):
    from app.bot.middlewares import permissions
    class Session:
        async def __aenter__(self): return self
        async def __aexit__(self, *args): pass
    monkeypatch.setattr(permissions, 'session_factory', Session)
    monkeypatch.setattr(permissions, 'require_permission', AsyncMock(side_effect=AccessDeniedError('Denied')))
    next_handler = AsyncMock()
    event = SimpleNamespace(answer=AsyncMock())
    await permissions.PermissionMiddleware('acts.manage')(next_handler, event,
        {'handler': SimpleNamespace(flags={}), 'system_user': SimpleNamespace(id=uuid.uuid4())})
    next_handler.assert_not_awaited()
    event.answer.assert_awaited_once()


async def test_filters_exclude_cancelled_rows_before_pagination():
    from app.db.repositories.available_work import for_act
    from sqlalchemy.dialects import postgresql
    session = SimpleNamespace(execute=AsyncMock(return_value=SimpleNamespace(all=lambda: [])))
    await for_act(session, SimpleNamespace(project_id=uuid.uuid4(), contract_id=uuid.uuid4()), limit=10, offset=10)
    query = str(session.execute.call_args.args[0].compile(dialect=postgresql.dialect()))
    assert 'act_lines.cancelled_at IS NULL' in query
    assert 'customer_acceptances.status =' in query
    assert query.index('WHERE work_entries.project_id') < query.index('LIMIT')


async def test_cancelled_acceptance_not_in_total_query():
    from app.db.repositories.customer_acceptance import total_accepted_volume
    session = SimpleNamespace(scalar=AsyncMock(return_value=Decimal('0')))
    await total_accepted_volume(session, uuid.uuid4())
    assert 'customer_acceptances.status =' in str(session.scalar.call_args.args[0])


async def test_finalized_act_rejects_line_cancellation(service_mocks, monkeypatch):
    monkeypatch.setattr(act.act_repo, 'get_for_update', AsyncMock(return_value=SimpleNamespace(status='finalized')))
    with pytest.raises(ValidationError):
        await act.remove_line(service_mocks, act_id=uuid.uuid4(), line_id=uuid.uuid4(), user_id=uuid.uuid4(), reason='fix')
    service_mocks.delete.assert_not_awaited()


def test_amount_overflow_is_validation_error_and_rounding_is_explicit():
    from app.core.numbers import line_amount
    with pytest.raises(ValidationError):
        line_amount(Decimal('999999999999999.999'), Decimal('9999999999999999.99'))
    assert line_amount(Decimal('0.001'), Decimal('15.00')) == Decimal('0.02')
    assert line_amount(Decimal('0.001'), Decimal('25.00')) == Decimal('0.02')


async def test_audit_entries_share_operation_request_id():
    session = SimpleNamespace(add=MagicMock(), flush=AsyncMock())
    @operation
    async def run():
        return [await audit.record(session, user_id=uuid.uuid4(), action='test', entity_type='act',
                                   entity_id='test', new_values={'amount': Decimal('1.20')}) for _ in range(2)]
    first, second = await run()
    assert first.request_id == second.request_id
    assert first.new_values['amount'] == '1.20'
