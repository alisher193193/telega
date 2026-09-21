from __future__ import annotations

import re
import tempfile
import os
from contextlib import contextmanager
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.acts import Act, ActLine
from app.models.directories import Contract, Project

_HEADERS = [
    "№",
    "Объект",
    "Договор/раздел",
    "Вид работы",
    "Место выполнения",
    "Ед. изм.",
    "Объём",
    "Цена",
    "Сумма",
    "Комментарий",
]

_MONEY_FORMAT = '#,##0.00" ₸"'
_MAX_COLUMN_WIDTH = 60


def _safe_filename(act_number: str) -> str:
    cleaned = re.sub(r"[^\w.\-]+", "_", act_number, flags=re.UNICODE).strip("_")
    return f"act_{cleaned or 'unnamed'}.xlsx"


def _autosize_columns(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        column_letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[column_letter].width = min(length + 2, _MAX_COLUMN_WIDTH)


def build_act_workbook(
    act: Act,
    lines: list[ActLine],
    project: Project,
    contract: Contract | None,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Акт"

    sheet.merge_cells("A1:J1")
    sheet["A1"] = f"Акт №{act.act_number} от {act.act_date.strftime('%d.%m.%Y')}"
    sheet["A1"].font = Font(bold=True, size=14)

    sheet.merge_cells("A2:J2")
    contract_label = contract.number if contract else "—"
    sheet["A2"] = f"Объект: {project.name}    Договор: {contract_label}"

    header_row = 4
    for col_index, header in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    row_index = header_row + 1
    for number, line in enumerate(lines, start=1):
        section_label = contract.number if contract else "—"
        sheet.cell(row=row_index, column=1, value=number)
        sheet.cell(row=row_index, column=2, value=project.name)
        sheet.cell(row=row_index, column=3, value=section_label)
        sheet.cell(row=row_index, column=4, value=line.work_type_name)
        sheet.cell(row=row_index, column=5, value=line.location_snapshot or "")
        sheet.cell(row=row_index, column=6, value=line.unit_name)
        sheet.cell(row=row_index, column=7, value=line.accepted_volume)

        price_cell = sheet.cell(row=row_index, column=8, value=line.unit_price)
        price_cell.number_format = _MONEY_FORMAT

        amount_cell = sheet.cell(row=row_index, column=9, value=line.amount)
        amount_cell.number_format = _MONEY_FORMAT

        comment_cell = sheet.cell(row=row_index, column=10, value=line.comment or "")
        comment_cell.alignment = Alignment(wrap_text=True)

        row_index += 1

    total_row = row_index
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=8)
    sheet.cell(row=total_row, column=1, value="Итого:").font = Font(bold=True)
    total_cell = sheet.cell(row=total_row, column=9, value=act.total_amount)
    total_cell.number_format = _MONEY_FORMAT
    total_cell.font = Font(bold=True)

    last_data_row = row_index - 1
    if last_data_row >= header_row + 1:
        sheet.auto_filter.ref = f"A{header_row}:J{last_data_row}"

    sheet.freeze_panes = f"A{header_row + 1}"

    _autosize_columns(sheet)

    # Explicit string cell types prevent openpyxl from interpreting untrusted text as formulas.
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, Decimal):
                # Excel numeric cells retain only 15 significant decimal digits.
                # Preserve large exact amounts as text instead of silently rounding.
                if len(cell.value.normalize().as_tuple().digits) > 15:
                    cell.value = format(cell.value, "f")
            if isinstance(cell.value, str):
                cell.data_type = "s"
    return workbook


def generate_act_excel(
    act: Act,
    lines: list[ActLine],
    project: Project,
    contract: Contract | None,
) -> Path:
    """Build the workbook and save it to a temporary file, returning its path.

    The caller is responsible for deleting the file after sending it to Telegram.
    """
    workbook = build_act_workbook(act, lines, project, contract)

    fd, name = tempfile.mkstemp(prefix="stage7_act_", suffix=".xlsx")
    os.close(fd)
    file_path = Path(name)
    try:
        workbook.save(file_path)
        return file_path
    except BaseException:
        file_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()


@contextmanager
def temporary_act_excel(act, lines, project, contract):
    path = generate_act_excel(act, lines, project, contract)
    try:
        yield path
    finally:
        path.unlink(missing_ok=True)
